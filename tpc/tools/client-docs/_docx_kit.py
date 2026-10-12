"""Generate the Tire Finder handoff as an editable Word document.

Mirrors the content of tire-finder-handoff.html, but adds a response box under
every open question so Jan can type answers directly into the file.
"""

import copy
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

# ── palette ──────────────────────────────────────────────────────────────────
INK        = RGBColor(0x22, 0x25, 0x29)
INK_SOFT   = RGBColor(0x5E, 0x62, 0x67)
ACCENT     = RGBColor(0xC0, 0x40, 0x40)
SETTLED    = RGBColor(0x2D, 0x6A, 0x4F)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

SH_HEAD    = "222529"   # table header fill
SH_ZEBRA   = "F4F4F5"   # alternating row
SH_SETTLED = "E8F4EE"   # settled callout
SH_QUESTION= "F6F6F7"   # question block
SH_RESPONSE= "FFFDF3"   # response box (warm = "type here")
SH_GAP     = "FDECEA"   # missing-data cell
RULE       = "DCDDE0"

SERIF = "Georgia"
SANS  = "Segoe UI"
MONO  = "Consolas"


# ── low-level helpers ────────────────────────────────────────────────────────
def shade(cell, fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(el)


def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def borders(cell, edges, color=RULE, sz=4):
    """edges: dict of edge -> True/False. Missing edges are cleared."""
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    node = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec:
            c, s = (spec if isinstance(spec, tuple) else (color, sz))
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(s))
            e.set(qn("w:color"), c)
        else:
            e.set(qn("w:val"), "nil")
        node.append(e)
    tcPr.append(node)


def no_borders(table):
    for row in table.rows:
        for cell in row.cells:
            borders(cell, {})


def keep_together(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def cant_split(table):
    """Stop individual rows breaking across a page."""
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def glue_rows(table, through=-1):
    """Keep every paragraph up to `through` with the next, so the rows of a
    block stay on one page together."""
    rows = table.rows[:through] if through != -1 else table.rows[:-1]
    for row in rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


def set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def run(paragraph, text, *, font=SANS, size=10.5, bold=False, italic=False,
        color=INK, caps=False, spacing=None):
    r = paragraph.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    if caps:
        r.font.all_caps = True
    if spacing is not None:
        rPr = r._element.get_or_add_rPr()
        el = OxmlElement("w:spacing")
        el.set(qn("w:val"), str(int(spacing * 20)))
        rPr.append(el)
    # make the font stick for non-ASCII too
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), font)
    return r


def para(container, text=None, *, space_before=0, space_after=6, line=1.3,
         align=None, **kw):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if text:
        run(p, text, **kw)
    return p


def rich(container, parts, *, space_before=0, space_after=6, line=1.3, **base):
    """parts: list of (text, overrides) tuples."""
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    for text, over in parts:
        kw = dict(base)
        kw.update(over)
        run(p, text, **kw)
    return p


def hrule(doc, color=RULE, sz=6, space_before=2, space_after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


# ── document scaffolding ─────────────────────────────────────────────────────
doc = Document()

sec = doc.sections[0]
sec.page_width  = Inches(8.5)
sec.page_height = Inches(11)
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = Inches(0.85)
sec.bottom_margin = Inches(0.9)

CONTENT_W = sec.page_width - sec.left_margin - sec.right_margin

normal = doc.styles["Normal"]
normal.font.name = SANS
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3

# Proof this document — and anything Jan types into it — as US English.
lang = OxmlElement("w:lang")
lang.set(qn("w:val"), "en-US")
lang.set(qn("w:eastAsia"), "en-US")
lang.set(qn("w:bidi"), "ar-SA")
normal.element.get_or_add_rPr().append(lang)

# footer with page numbers
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(footer_p, "Rene Herse Cycles · Tire Finder · page ", font=MONO, size=8, color=INK_SOFT)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer_p._p.append(fld)


def h2(text):
    p = para(doc, text, font=SERIF, size=17, bold=True, color=INK,
             space_before=20, space_after=2, line=1.2)
    keep_together(p)
    # The rule under a heading must also stay with what follows, or a heading can
    # sit alone at the foot of a page with its section overleaf.
    keep_together(hrule(doc, space_after=9))
    return p


def h3(text):
    p = para(doc, text, font=SANS, size=9, bold=True, color=INK_SOFT,
             caps=True, spacing=1.0, space_before=14, space_after=5)
    keep_together(p)
    return p


def body(text, space_after=7):
    return para(doc, text, size=10.5, color=INK, space_after=space_after)


def boxed_table(fill, border_color=None, border_edge="left", border_sz=18):
    """Single-cell table used as a callout box."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_col_widths(t, [CONTENT_W])
    cell = t.cell(0, 0)
    cell.text = ""
    shade(cell, fill)
    cell_margins(cell, top=140, bottom=140, left=180, right=180)
    edges = {}
    if border_color:
        edges[border_edge] = (border_color, border_sz)
    borders(cell, edges)
    return cell


def spacer(pts=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    run(p, "", size=pts / 2)
    return p


def data_table(headers, rows, widths, *, gap_cells=frozenset(), mono_cols=frozenset(),
               compact=False, vpad=70):
    """headers: list[str]; rows: list[list[str]]; gap_cells: set of (row, col).

    compact=True tightens the horizontal padding so short column labels fit on
    one line in a narrow column.
    """
    pad = 55 if compact else 120
    head_size = 7.5 if compact else 8
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_col_widths(t, widths)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    for j, head in enumerate(headers):
        cell = t.cell(0, j)
        cell.text = ""
        shade(cell, SH_HEAD)
        cell_margins(cell, top=vpad, bottom=vpad, left=pad, right=pad)
        borders(cell, {})
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        run(p, head, font=MONO, size=head_size, bold=True, color=WHITE, caps=True,
            spacing=0.2 if compact else 0.4)

    for i, cols in enumerate(rows):
        for j, text in enumerate(cols):
            cell = t.cell(i + 1, j)
            cell.text = ""
            is_gap = (i, j) in gap_cells
            if is_gap:
                shade(cell, SH_GAP)
            elif i % 2 == 1:
                shade(cell, SH_ZEBRA)
            cell_margins(cell, top=vpad, bottom=vpad, left=pad, right=pad)
            borders(cell, {"bottom": (RULE, 4)})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.25
            if is_gap:
                run(p, text, font=MONO, size=9, bold=True, color=ACCENT)
            elif j in mono_cols:
                run(p, text, font=MONO, size=8.5, color=INK)
            else:
                # allow inline bold via **markers**
                for k, chunk in enumerate(text.split("**")):
                    if chunk:
                        run(p, chunk, size=9.5, bold=bool(k % 2), color=INK)
    cant_split(t)
    return t


def question(qid, ask, why):
    """Question block with an attached response box for Jan."""
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    id_w = Inches(0.5)
    set_col_widths(t, [id_w, CONTENT_W - id_w])

    # ── row 0: the question
    idc = t.cell(0, 0)
    idc.text = ""
    shade(idc, SH_QUESTION)
    cell_margins(idc, top=130, bottom=110, left=150, right=40)
    borders(idc, {"left": (ACCENT.__str__(), 18)})
    p = idc.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run(p, qid, font=MONO, size=10, bold=True, color=ACCENT)

    bodyc = t.cell(0, 1)
    bodyc.text = ""
    shade(bodyc, SH_QUESTION)
    cell_margins(bodyc, top=130, bottom=110, left=40, right=150)
    borders(bodyc, {})
    p = bodyc.paragraphs[0]
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.3
    for k, chunk in enumerate(ask.split("**")):
        if chunk:
            run(p, chunk, size=10.5, bold=True, color=INK)
    p2 = bodyc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.3
    for k, chunk in enumerate(why.split("**")):
        if chunk:
            run(p2, chunk, size=9.5, color=INK_SOFT,
                font=MONO if k % 2 else SANS)

    # ── row 1: response box
    lab = t.cell(1, 0)
    lab.text = ""
    shade(lab, SH_RESPONSE)
    cell_margins(lab, top=110, bottom=140, left=150, right=40)
    borders(lab, {"left": (ACCENT.__str__(), 18), "top": ("E8E4D8", 4)})
    p = lab.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run(p, "→", font=SANS, size=10, bold=True, color=INK_SOFT)

    ans = t.cell(1, 1)
    ans.text = ""
    shade(ans, SH_RESPONSE)
    cell_margins(ans, top=110, bottom=140, left=40, right=150)
    borders(ans, {"top": ("E8E4D8", 4)})
    p = ans.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.35
    run(p, "Jan: ", font=SANS, size=9.5, bold=True, color=INK_SOFT)
    run(p, "[type your answer here]", font=SANS, size=9.5, italic=True,
        color=RGBColor(0xA8, 0xA2, 0x8E))

    # a question must never be separated from the box it is answered in
    cant_split(t)
    glue_rows(t)

    spacer(7)
    return t


